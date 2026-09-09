from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

from pai_auto.importer import import_workbook, select_batch


def _book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vacunación SIS"
    for col in range(1, 223):
        ws.cell(15, col, f"C{col}")
    ws["B16"] = date(2026, 9, 8)
    ws["C16"] = "CC"
    ws["D16"] = "123456"
    ws["HK16"] = "VACUNADOR DEMO"
    # BCG BS:BW => dosis, lote, jeringa, diluyente, observación
    ws.cell(16, column_index_from_string("BS"), "Única")
    ws.cell(16, column_index_from_string("BT"), "LOTE-BCG-1")
    ws.cell(16, column_index_from_string("BU"), "JER-1")
    wb.save(path)


def test_imports_one_application_from_222_contract(tmp_path):
    path = tmp_path / "registro.xlsx"
    _book(path)
    result = import_workbook(path)
    assert result.total_source_rows == 1
    assert len(result.records) == 1
    record = result.records[0]
    assert record.source_row == 16
    assert record.document_number == "123456"
    assert record.biologic_code == "PAI_BCG"
    assert record.dose == "Única"
    assert record.lot == "LOTE-BCG-1"


def test_batch_sizes_are_closed(tmp_path):
    path = tmp_path / "registro.xlsx"
    _book(path)
    records = import_workbook(path).records
    assert len(select_batch(records, start_source_row=16, batch_size=1)) == 1
    try:
        select_batch(records, start_source_row=16, batch_size=10)
    except ValueError as exc:
        assert "Permitidos" in str(exc)
    else:
        raise AssertionError("10 no debe ser un lote permitido")
