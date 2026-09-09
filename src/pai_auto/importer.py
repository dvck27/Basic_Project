from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .models import ApplicationRecord, ImportIssue, ImportResult


@dataclass(frozen=True, slots=True)
class Group:
    code: str
    name: str
    start: str
    fields: tuple[str, ...]


# Contrato PAI BN:FF del Registro Diario SIS v1 del repositorio original.
PAI_GROUPS: tuple[Group, ...] = (
    Group("PAI_COVID19", "COVID-19", "BN", ("dose", "laboratory", "lot", "syringe", "diluent")),
    Group("PAI_BCG", "BCG", "BS", ("dose", "lot", "syringe", "diluent", "observation")),
    Group("PAI_HEPATITIS_B", "Hepatitis B", "BX", ("dose", "lot", "syringe", "observation")),
    Group("PAI_POLIO_INACTIVADO", "Polio Inactivado (Vacuna inyectable)", "CB", ("dose", "lot", "syringe")),
    Group("PAI_PENTAVALENTE", "Pentavalente", "CE", ("dose", "lot", "syringe")),
    Group("PAI_HEXAVALENTE", "Hexavalente", "CH", ("dose", "lot", "syringe")),
    Group("PAI_DPT", "Difteria, Tos ferina y Tétanos - DPT", "CK", ("dose", "lot", "syringe")),
    Group("PAI_DTPA_PEDIATRICO", "DTPa pediátrico", "CN", ("dose", "lot", "syringe")),
    Group("PAI_TD_PEDIATRICO", "TD pediátrico", "CQ", ("dose", "lot", "syringe")),
    Group("PAI_ROTAVIRUS", "Rotavirus (vacuna oral)", "CT", ("dose", "lot")),
    Group("PAI_NEUMOCOCO", "Neumococo", "CV", ("type", "dose", "lot", "syringe")),
    Group("PAI_TRIPLE_VIRAL_SRP", "Triple viral - SRP", "CZ", ("dose", "lot", "syringe", "diluent")),
    Group("PAI_SR_MULTIDOSIS", "Sarampión - Rubeola - SR Multidosis", "DD", ("dose", "lot", "syringe", "diluent")),
    Group("PAI_FIEBRE_AMARILLA", "Fiebre amarilla", "DH", ("dose", "lot", "syringe", "diluent")),
    Group("PAI_HEPATITIS_A_PEDIATRICA", "Hepatitis A pediátrica", "DL", ("dose", "lot", "syringe")),
    Group("PAI_VARICELA", "Varicela", "DO", ("dose", "lot", "syringe", "diluent")),
    Group("PAI_TD_ADULTO", "Toxoide tetánico y diftérico de Adulto", "DS", ("dose", "lot", "syringe")),
    Group("PAI_DTPA_ADULTO", "dTpa adulto", "DV", ("dose", "lot", "syringe")),
    Group("PAI_VSR", "Virus Sincitial Respiratorio", "DY", ("dose", "lot", "syringe")),
    Group("PAI_INFLUENZA", "Influenza", "EB", ("dose", "lot", "syringe", "observation")),
    Group("PAI_DENGUE", "DENGUE", "EF", ("dose", "lot", "syringe")),
    Group("PAI_VPH", "VPH", "EI", ("dose", "lot", "syringe")),
    Group("PAI_ANTIRRABICA_HUMANA", "Antirrábica  Humana (vacuna)", "EL", ("dose", "lot", "syringe", "diluent", "observation")),
    Group("PAI_ANTIRRABICO_HUMANO_SUERO", "Antirrábico Humano (suero)", "EQ", ("dose", "lot")),
    Group("PAI_INMUNOGLOBULINA_HEPATITIS_B", "Hepatitis B (Inmunoglobulina)", "ES", ("dose", "lot", "syringe", "observation")),
    Group("PAI_INMUNOGLOBULINA_ANTITETANICA", "INMUNOGLOBULINA ANTI TETANICA (Suero homólogo)", "EW", ("dose", "lot", "syringe")),
    Group("PAI_ANTITOXINA_DIFTERICA", "Antitoxina difterica (Suero heterólogo)", "EZ", ("dose", "lot", "syringe")),
    Group("PAI_MENINGOCOCO_ACWY", "Meningococo  de los serogrupos A, C, W-135 e Y", "FC", ("dose", "lot", "syringe", "diluent")),
)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).strip().split())


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _group_values(row: tuple[object, ...], group: Group) -> dict[str, str]:
    start = column_index_from_string(group.start) - 1
    values = row[start : start + len(group.fields)]
    return {name: _text(value) for name, value in zip(group.fields, values, strict=True)}


def import_workbook(path: str | Path, *, start_row: int = 16) -> ImportResult:
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Vacunación SIS" not in workbook.sheetnames:
            raise ValueError("No existe la hoja requerida 'Vacunación SIS'.")
        sheet = workbook["Vacunación SIS"]
        if sheet.max_column != 222:
            raise ValueError(f"Contrato desconocido: se esperaban 222 columnas y hay {sheet.max_column}.")

        records: list[ApplicationRecord] = []
        issues: list[ImportIssue] = []
        total_rows = 0
        for index, row in enumerate(
            sheet.iter_rows(min_row=start_row, max_col=222, values_only=True), start=start_row
        ):
            if not any(value not in (None, "") for value in row):
                continue
            total_rows += 1
            document_type = _text(row[2])
            document_number = _text(row[3])
            vaccination_date = _date(row[1])
            vaccinator = _text(row[column_index_from_string("HK") - 1])
            row_has_application = False

            if not document_number or vaccination_date is None:
                issues.append(
                    ImportIssue(
                        index,
                        "FILA_SIN_IDENTIDAD_O_FECHA",
                        "Falta documento o fecha de atención válida.",
                    )
                )
                continue

            for group in PAI_GROUPS:
                values = _group_values(row, group)
                if not any(values.values()):
                    continue
                row_has_application = True
                dose = values.get("dose", "")
                lot = values.get("lot", "")
                if not dose:
                    issues.append(
                        ImportIssue(
                            index,
                            "GRUPO_SIN_DOSIS",
                            f"{group.name}: hay datos pero falta dosis/frascos.",
                        )
                    )
                    continue
                if not lot:
                    issues.append(ImportIssue(index, "DOSIS_SIN_LOTE", f"{group.name}: falta lote."))
                    continue
                records.append(
                    ApplicationRecord(
                        source_row=index,
                        document_type=document_type,
                        document_number=document_number,
                        vaccination_date=vaccination_date,
                        vaccinator=vaccinator,
                        biologic_code=group.code,
                        biologic_name=group.name,
                        dose=dose,
                        lot=lot,
                        syringe_lot=values.get("syringe", ""),
                        diluent_lot=values.get("diluent", ""),
                        biologic_type=values.get("type", ""),
                        observation=values.get("observation", ""),
                    )
                )
            if not row_has_application:
                issues.append(
                    ImportIssue(index, "FILA_SIN_APLICACION", "La fila no contiene aplicaciones PAI.")
                )

        return ImportResult(path.name, total_rows, records, issues)
    finally:
        workbook.close()


def select_batch(
    records: Iterable[ApplicationRecord], *, start_source_row: int, batch_size: int
) -> list[ApplicationRecord]:
    allowed = {1, 20, 50, 100}
    if batch_size not in allowed:
        raise ValueError(f"Lote inválido: {batch_size}. Permitidos: {sorted(allowed)}")
    eligible = [record for record in records if record.source_row >= start_source_row]
    selected_rows: list[int] = []
    seen: set[int] = set()
    for record in eligible:
        if record.source_row not in seen:
            seen.add(record.source_row)
            selected_rows.append(record.source_row)
            if len(selected_rows) == batch_size:
                break
    selected = set(selected_rows)
    return [record for record in eligible if record.source_row in selected]
